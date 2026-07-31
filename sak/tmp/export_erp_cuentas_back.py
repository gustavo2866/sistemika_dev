from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import psycopg

ROOT = Path(r"d:/gpalmieri/proyectos/sistemika_dev/sak")
ENV_PATH = ROOT / "backend" / ".env"
OUTPUT_PATH = ROOT / "backend" / "data" / "erp_cuentas_back.xlsx"


def get_db_url() -> str:
    raw = ENV_PATH.read_text(encoding="utf-8")
    match = re.search(r"(?m)^DATABASE_URL\s*=\s*(.+)\s*$", raw)
    if not match:
        raise RuntimeError("DATABASE_URL no encontrado en backend/.env")
    return match.group(1).strip().strip('"').strip("'").replace(
        "postgresql+psycopg://", "postgresql://", 1
    )


def normalize_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    if isinstance(value, date):
        return value
    if isinstance(value, Decimal):
        return value
    return value


def build_workbook(rows: list[tuple[Any, ...]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "erp_cuentas_back"

    headers = [
        "id",
        "rubro_id",
        "rubro_nombre",
        "nro_cuenta",
        "cod_cuenta",
        "descripcion",
        "activo",
        "created_at",
        "updated_at",
        "deleted_at",
        "version",
        "proyectos_concepto_id",
        "concepto_nombre",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="E5E7EB")
    header_font = Font(bold=True, color="111827")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([normalize_value(value) for value in row])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = '#,##0.00'
            elif isinstance(cell.value, (datetime, date)):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss' if isinstance(cell.value, datetime) else 'yyyy-mm-dd'

    widths = {
        "A": 10,
        "B": 12,
        "C": 30,
        "D": 12,
        "E": 16,
        "F": 48,
        "G": 10,
        "H": 20,
        "I": 20,
        "J": 16,
        "K": 12,
        "L": 18,
        "M": 26,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return wb


def main() -> int:
    db_url = get_db_url()
    query = """
        select
            ec.id,
            ec.rubro_id,
            er.nombre as rubro_nombre,
            ec.nro_cuenta,
            ec.cod_cuenta,
            ec.descripcion,
            ec.activo,
            ec.created_at,
            ec.updated_at,
            ec.deleted_at,
            ec.version,
            ec.proyectos_concepto_id,
            pc.nombre as concepto_nombre
        from public.erp_cuentas ec
        left join public.erp_rubros er on er.id = ec.rubro_id
        left join public.proyectos_conceptos pc on pc.id = ec.proyectos_concepto_id
        order by ec.nro_cuenta, ec.cod_cuenta, ec.id
    """

    with psycopg.connect(db_url) as conn:
        with conn.cursor() as cur:
            cur.execute(query)
            rows = cur.fetchall()

    wb = build_workbook(rows)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"exported_rows={len(rows)}")
    print(f"output_path={OUTPUT_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
