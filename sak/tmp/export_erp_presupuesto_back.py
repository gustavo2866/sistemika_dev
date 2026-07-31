from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import psycopg


ROOT = Path(r"d:/gpalmieri/proyectos/sistemika_dev/sak")
ENV_PATH = ROOT / "backend" / ".env"
OUTPUT_PATH = ROOT / "backend" / "data" / "erp_presupuesto_back.xlsx"


def get_db_url() -> str:
    raw = ENV_PATH.read_text(encoding="utf-8")
    match = re.search(r"(?m)^DATABASE_URL\s*=\s*(.+)\s*$", raw)
    if not match:
        raise RuntimeError("DATABASE_URL no encontrado en backend/.env")
    return match.group(1).strip().strip('"').strip("'").replace(
        "postgresql+psycopg://", "postgresql://", 1
    )


def normalize_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    if isinstance(value, date):
        return value
    if isinstance(value, Decimal):
        return value
    return value


def build_workbook(rows: list[tuple[Any, ...]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "erp_presupuesto_back"

    headers = [
        "id",
        "fecha",
        "proyecto_id",
        "proyecto_nombre",
        "erp_cuenta_id",
        "nro_cuenta",
        "codigo_cuenta",
        "nombre",
        "egreso",
        "ingres",
        "obreros_cantidad",
        "obreros_costo",
        "created_at",
        "updated_at",
        "deleted_at",
        "version",
        "real_egreso",
        "real_ingreso",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="E5E7EB")
    header_font = Font(bold=True, color="111827")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([normalize_value(value) for value in row])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = '#,##0.00'
            elif isinstance(cell.value, (datetime, date)):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss' if isinstance(cell.value, datetime) else 'yyyy-mm-dd'

    widths = {
        "A": 10,
        "B": 14,
        "C": 12,
        "D": 36,
        "E": 12,
        "F": 12,
        "G": 14,
        "H": 44,
        "I": 16,
        "J": 16,
        "K": 18,
        "L": 16,
        "M": 20,
        "N": 20,
        "O": 16,
        "P": 12,
        "Q": 16,
        "R": 16,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return wb


def main() -> int:
    db_url = get_db_url()
    query = """
        select
            ep.id,
            ep.fecha,
            ep.proyecto_id,
            p.nombre as proyecto_nombre,
            ep.erp_cuenta_id,
            ec.nro_cuenta,
            ec.cod_cuenta as codigo_cuenta,
            ec.descripcion as nombre,
            ep.egreso,
            ep.ingres,
            ep.obreros_cantidad,
            ep.obreros_costo,
            ep.created_at,
            ep.updated_at,
            ep.deleted_at,
            ep.version,
            ep.real_egreso,
            ep.real_ingreso
        from public.erp_presupuestos ep
        left join public.proyectos p on p.id = ep.proyecto_id
        left join public.erp_cuentas ec on ec.id = ep.erp_cuenta_id
        order by p.nombre nulls last, ep.fecha, ec.nro_cuenta nulls last, ec.cod_cuenta nulls last, ep.id
    """

    with psycopg.connect(db_url) as conn:
        with conn.cursor() as cur:
            cur.execute(query)
            rows = cur.fetchall()

    wb = build_workbook(rows)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"exported_rows={len(rows)}")
    print(f"output_path={OUTPUT_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())