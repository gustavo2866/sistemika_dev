from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook
import psycopg

ROOT = Path(r"d:/gpalmieri/proyectos/sistemika_dev/sak")
ENV_PATH = ROOT / "backend" / ".env"
SOURCE_PATH = ROOT / "backend" / "data" / "erp_presupuesto_01.xls.xlsx"

EXCLUDED_PROJECT_NAMES = {"catamarca y corrientes"}
PROJECT_NAME_MAP = {
    "axion": 18,
    "francia": 19,
    "rioja": 10,
    "torre sp": 14,
}


@dataclass(frozen=True)
class BudgetKey:
    proyecto_id: int
    erp_cuenta_id: int
    fecha: datetime


def get_db_url() -> str:
    raw = ENV_PATH.read_text(encoding="utf-8")
    match = re.search(r"(?m)^DATABASE_URL\s*=\s*(.+)\s*$", raw)
    if not match:
        raise RuntimeError("DATABASE_URL no encontrado en backend/.env")
    return match.group(1).strip().strip('"').strip("'").replace(
        "postgresql+psycopg://", "postgresql://", 1
    )


def normalize_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def parse_decimal(value: Any) -> Decimal:
    text = normalize_text(value).replace("$", "").replace(" ", "")
    if not text:
        return Decimal("0")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text and "." not in text:
        text = text.replace(",", ".")
    text = re.sub(r"[^0-9\.-]", "", text)
    if text in {"", "-", ".", "-."}:
        return Decimal("0")
    try:
        return Decimal(text)
    except Exception:
        return Decimal("0")


def parse_period(value: Any) -> datetime:
    text = normalize_text(value)
    if not text:
        raise ValueError("periodo vacío")

    if re.fullmatch(r"\d+(\.\d+)?", text):
        base = datetime(1899, 12, 30)
        return base.fromordinal(base.toordinal() + int(float(text)))

    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            parsed = datetime.strptime(text[:10], fmt)
            return datetime(parsed.year, parsed.month, 1)
        except Exception:
            pass

    match = re.search(r"(20\d{2})[-/](\d{1,2})", text)
    if match:
        return datetime(int(match.group(1)), int(match.group(2)), 1)

    match = re.search(r"(\d{1,2})[-/](20\d{2})", text)
    if match:
        return datetime(int(match.group(2)), int(match.group(1)), 1)

    raise ValueError(f"No se pudo normalizar el periodo: {value!r}")


def normalize_obra(value: Any) -> str:
    return normalize_text(value)


def load_source_rows() -> tuple[list[dict[str, Any]], set[datetime]]:
    workbook = load_workbook(SOURCE_PATH, data_only=True)
    sheet = workbook.active
    headers = [normalize_text(sheet.cell(row=1, column=col).value) for col in range(1, sheet.max_column + 1)]
    header_map = {header: index for index, header in enumerate(headers)}

    required = ["obra", "mov", "cuenta", "periodo", "importe", "empleados"]
    missing = [name for name in required if name not in header_map]
    if missing:
        raise RuntimeError(f"Faltan columnas en el archivo fuente: {missing}")

    source_rows: list[dict[str, Any]] = []
    source_periods: set[datetime] = set()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        obra = normalize_obra(row[header_map["obra"]])
        if obra in EXCLUDED_PROJECT_NAMES:
            continue

        mov = normalize_text(row[header_map["mov"]]).upper()
        if mov not in {"EG", "ING"}:
            continue

        cuenta_raw = row[header_map["cuenta"]]
        try:
            erp_cuenta_num = int(float(str(cuenta_raw).replace(",", ".").strip()))
        except Exception as exc:
            raise RuntimeError(f"Cuenta inválida en fila {row!r}") from exc

        if obra not in PROJECT_NAME_MAP:
            raise RuntimeError(f"Obra no mapeada: {obra!r}")

        fecha = parse_period(row[header_map["periodo"]])
        importe = parse_decimal(row[header_map["importe"]])
        empleados = parse_decimal(row[header_map["empleados"]])

        source_periods.add(fecha)
        source_rows.append(
            {
                "obra": obra,
                "mov": mov,
                "proyecto_id": PROJECT_NAME_MAP[obra],
                "erp_cuenta_num": erp_cuenta_num,
                "fecha": fecha,
                "importe": importe,
                "empleados": empleados,
            }
        )

    return source_rows, source_periods


def resolve_project_ids(conn: psycopg.Connection) -> dict[int, str]:
    with conn.cursor() as cur:
        cur.execute(
            """
            select id, nombre
            from public.proyectos
            where lower(nombre) not in ('catamarca y corrientes')
            """
        )
        return {int(project_id): str(name) for project_id, name in cur.fetchall()}


def resolve_cuenta_ids(conn: psycopg.Connection) -> dict[int, int]:
    with conn.cursor() as cur:
        cur.execute("select id, nro_cuenta from public.erp_cuentas where nro_cuenta is not null")
        return {int(nro_cuenta): int(cuenta_id) for cuenta_id, nro_cuenta in cur.fetchall()}


def cleanup_existing_rows(conn: psycopg.Connection, project_ids: list[int], periods: set[datetime]) -> int:
    if not project_ids or not periods:
        return 0

    with conn.cursor() as cur:
        cur.execute(
            """
            delete from public.erp_presupuestos
            where proyecto_id = any(%s)
              and fecha = any(%s)
            """,
            (project_ids, sorted({period.date() for period in periods})),
        )
        return cur.rowcount or 0


def main() -> int:
    db_url = get_db_url()
    source_rows, source_periods = load_source_rows()

    eg_rows: dict[BudgetKey, dict[str, Decimal]] = defaultdict(lambda: {"egreso": Decimal("0"), "empleados": Decimal("0")})
    ing_rows: dict[BudgetKey, Decimal] = defaultdict(lambda: Decimal("0"))

    for row in source_rows:
        key = BudgetKey(row["proyecto_id"], row["erp_cuenta_num"], row["fecha"])
        if row["mov"] == "EG":
            eg_rows[key]["egreso"] += row["importe"]
            eg_rows[key]["empleados"] += row["empleados"]
        elif row["mov"] == "ING":
            ing_rows[key] += row["importe"]

    with psycopg.connect(db_url) as conn:
        conn.autocommit = False
        try:
            project_ids = sorted(set(PROJECT_NAME_MAP.values()))
            deleted = cleanup_existing_rows(conn, project_ids, source_periods)
            cuenta_ids = resolve_cuenta_ids(conn)

            inserted_ids: dict[BudgetKey, int] = {}
            eg_inserted = 0
            eg_no_match: list[BudgetKey] = []

            with conn.cursor() as cur:
                for key, values in eg_rows.items():
                    cuenta_id = cuenta_ids.get(key.erp_cuenta_id)
                    if not cuenta_id:
                        eg_no_match.append(key)
                        continue

                    cur.execute(
                        """
                        insert into public.erp_presupuestos (
                            fecha,
                            proyecto_id,
                            erp_cuenta_id,
                            egreso,
                            ingres,
                            real_egreso,
                            real_ingreso,
                            obreros_cantidad,
                            obreros_costo,
                            version,
                            created_at,
                            updated_at
                        )
                        values (
                            %s, %s, %s,
                            %s, 0,
                            0, 0,
                            %s, 0,
                            1,
                            now(),
                            now()
                        )
                        returning id
                        """,
                        (
                            key.fecha.date(),
                            key.proyecto_id,
                            cuenta_id,
                            values["egreso"],
                            values["empleados"],
                        ),
                    )
                    inserted_ids[key] = int(cur.fetchone()[0])
                    eg_inserted += 1

                ing_updated = 0
                ing_inserted = 0
                ing_no_match: list[BudgetKey] = []
                for key, importe in ing_rows.items():
                    presupuesto_id = inserted_ids.get(key)
                    if not presupuesto_id:
                        cuenta_id = cuenta_ids.get(key.erp_cuenta_id)
                        if not cuenta_id:
                            ing_no_match.append(key)
                            continue

                        cur.execute(
                            """
                            insert into public.erp_presupuestos (
                                fecha,
                                proyecto_id,
                                erp_cuenta_id,
                                egreso,
                                ingres,
                                real_egreso,
                                real_ingreso,
                                obreros_cantidad,
                                obreros_costo,
                                version,
                                created_at,
                                updated_at
                            )
                            values (
                                %s, %s, %s,
                                0, %s,
                                0, 0,
                                0, 0,
                                1,
                                now(),
                                now()
                            )
                            returning id
                            """,
                            (
                                key.fecha.date(),
                                key.proyecto_id,
                                cuenta_id,
                                importe,
                            ),
                        )
                        inserted_ids[key] = int(cur.fetchone()[0])
                        ing_inserted += 1
                        continue

                    cur.execute(
                        """
                        update public.erp_presupuestos
                        set ingres = %s,
                            updated_at = now(),
                            version = coalesce(version, 1) + 1
                        where id = %s
                        """,
                        (importe, presupuesto_id),
                    )
                    ing_updated += cur.rowcount or 0

            conn.commit()

            print(f"registros_eliminados={deleted}")
            print(f"eg_insertados={eg_inserted}")
            print(f"ing_insertados={ing_inserted}")
            print(f"ing_actualizados={ing_updated}")
            print(f"eg_no_match={len(eg_no_match)}")
            print(f"ing_no_match={len(ing_no_match)}")
            if eg_no_match:
                print("eg_no_match_keys=")
                for key in eg_no_match:
                    print(key)
            if ing_no_match:
                print("ing_no_match_keys=")
                for key in ing_no_match:
                    print(key)
            return 0
        except Exception:
            conn.rollback()
            raise


if __name__ == "__main__":
    raise SystemExit(main())
