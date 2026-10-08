from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlmodel import select

from app.models.erp.cash_budget_excel import (
    BUDGET_SHEET_NAME,
    cash_budget_workbook_bytes,
    import_cash_budget_workbook,
)
from app.models.erp.cash_cuenta import ErpCashCuenta
from app.models.erp.cash_proyectado import ErpCashProyectado


def _seed_accounts(db_session):
    income = ErpCashCuenta(descripcion="INGR x VENTAS", tipo="Ingreso")
    expense = ErpCashCuenta(descripcion="EGR x COMPRAS", tipo="Egreso")
    fund = ErpCashCuenta(descripcion="FONDO", tipo="Fondo")
    db_session.add_all([income, expense, fund])
    db_session.commit()
    db_session.refresh(income)
    db_session.refresh(expense)
    return income, expense


def test_export_cash_budget_builds_protected_24_month_template(db_session):
    income, expense = _seed_accounts(db_session)
    db_session.add_all(
        [
            ErpCashProyectado(
                cuenta_cash_id=income.id,
                fecha_periodo=date(2026, 1, 1),
                tipo="PROYECCION",
                importe=Decimal("150000000"),
            ),
            ErpCashProyectado(
                cuenta_cash_id=expense.id,
                fecha_periodo=date(2026, 1, 1),
                tipo="PROYECCION",
                importe=Decimal("-40000000"),
            ),
        ]
    )
    db_session.commit()

    workbook = load_workbook(BytesIO(cash_budget_workbook_bytes(db_session, 2026)), data_only=False)
    worksheet = workbook[BUDGET_SHEET_NAME]

    assert worksheet.max_column == 26
    assert worksheet["A1"].value == "ID"
    assert worksheet["B1"].value == "Cuenta Cash"
    assert worksheet["C1"].value.date() == date(2026, 1, 1)
    assert worksheet["Z1"].value.date() == date(2027, 12, 1)
    assert worksheet["A2"].protection.locked is True
    assert worksheet["B2"].protection.locked is True
    assert worksheet["C2"].protection.locked is False
    assert worksheet["C2"].value == 150
    assert worksheet["C3"].value == -40
    assert worksheet["C2"].number_format == '#,##0.00 "M";[Red](#,##0.00 "M");-'
    assert worksheet["B5"].value == "TOTAL INGRESOS"
    assert worksheet["B6"].value == "TOTAL EGRESOS"
    assert worksheet["B7"].value == "SALDO DEL PERIODO"
    assert worksheet["C5"].value == "=SUM(C2:C2)"
    assert worksheet["C6"].value == "=SUM(C3:C3)"
    assert worksheet["C7"].value == "=C5+C6"
    assert worksheet.protection.sheet is True
    assert worksheet.protection.formatColumns is False


def test_import_cash_budget_updates_projection_values(db_session):
    income, expense = _seed_accounts(db_session)
    content = cash_budget_workbook_bytes(db_session, 2026)
    workbook = load_workbook(BytesIO(content))
    worksheet = workbook[BUDGET_SHEET_NAME]
    worksheet["C2"] = 1250.75
    worksheet["C3"] = -475.25
    output = BytesIO()
    workbook.save(output)

    result = import_cash_budget_workbook(db_session, output.getvalue(), 2026)

    assert result["created"] == 48
    rows = db_session.exec(
        select(ErpCashProyectado).where(
            ErpCashProyectado.fecha_periodo == date(2026, 1, 1),
            ErpCashProyectado.tipo == "PROYECCION",
        )
    ).all()
    values = {row.cuenta_cash_id: row.importe for row in rows}
    assert values[income.id] == Decimal("1250750000.00")
    assert values[expense.id] == Decimal("-475250000.00")


def test_import_cash_budget_rejects_modified_account_name(db_session):
    _seed_accounts(db_session)
    workbook = load_workbook(BytesIO(cash_budget_workbook_bytes(db_session, 2026)))
    workbook[BUDGET_SHEET_NAME]["B2"] = "Cuenta modificada"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ValueError, match="nombre de la cuenta Cash"):
        import_cash_budget_workbook(db_session, output.getvalue(), 2026)
