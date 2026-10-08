from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from math import isfinite

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select
from sqlalchemy import func

from app.models.base import current_utc_time
from app.models.erp.cash_cuenta import ErpCashCuenta
from app.models.erp.cash_proyectado import ErpCashProyectado


BUDGET_TYPE = "PROYECCION"
BUDGET_MONTHS = 24
BUDGET_SHEET_NAME = "Presupuesto Cash"
BUDGET_TOTAL_LABELS = ("TOTAL INGRESOS", "TOTAL EGRESOS", "SALDO DEL PERIODO")
MILLION = Decimal("1000000")


def budget_months(year: int) -> tuple[date, ...]:
    start_index = year * 12
    return tuple(
        date((start_index + offset) // 12, (start_index + offset) % 12 + 1, 1)
        for offset in range(BUDGET_MONTHS)
    )


def _cash_accounts(session: Session) -> list[ErpCashCuenta]:
    accounts = session.exec(
        select(ErpCashCuenta).where(
            ErpCashCuenta.deleted_at.is_(None),
            func.upper(func.trim(ErpCashCuenta.tipo)).in_(("INGRESO", "EGRESO")),
        )
    ).all()
    return sorted(
        accounts,
        key=lambda account: (
            0 if str(account.tipo or "").strip().upper() == "INGRESO" else 1,
            str(account.descripcion).upper(),
            int(account.id or 0),
        ),
    )


def build_cash_budget_workbook(session: Session, year: int) -> Workbook:
    months = budget_months(year)
    accounts = _cash_accounts(session)
    account_ids = [int(account.id) for account in accounts if account.id is not None]
    projected_rows = []
    if account_ids:
        projected_rows = session.exec(
            select(ErpCashProyectado).where(
                ErpCashProyectado.deleted_at.is_(None),
                ErpCashProyectado.tipo == BUDGET_TYPE,
                ErpCashProyectado.cuenta_cash_id.in_(account_ids),
                ErpCashProyectado.fecha_periodo >= months[0],
                ErpCashProyectado.fecha_periodo <= months[-1],
            )
        ).all()
    values_by_key = {
        (int(row.cuenta_cash_id), row.fecha_periodo): Decimal(row.importe or 0)
        for row in projected_rows
    }

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = BUDGET_SHEET_NAME
    worksheet.sheet_view.showGridLines = False

    headers: list[object] = ["ID", "Cuenta Cash", *months]
    worksheet.append(headers)
    for column in range(3, 3 + BUDGET_MONTHS):
        worksheet.cell(1, column).number_format = "mmm-yy"

    for account in accounts:
        worksheet.append(
            [
                int(account.id),
                account.descripcion,
                *[
                    float(
                        values_by_key.get((int(account.id), month), Decimal("0"))
                        / MILLION
                    )
                    for month in months
                ],
            ]
        )

    first_data_row = 2
    last_data_row = len(accounts) + 1
    income_rows = [
        index + first_data_row
        for index, account in enumerate(accounts)
        if str(account.tipo or "").strip().upper() == "INGRESO"
    ]
    expense_rows = [
        index + first_data_row
        for index, account in enumerate(accounts)
        if str(account.tipo or "").strip().upper() == "EGRESO"
    ]
    total_income_row = last_data_row + 2
    total_expense_row = total_income_row + 1
    period_balance_row = total_expense_row + 1
    worksheet.cell(total_income_row, 2, BUDGET_TOTAL_LABELS[0])
    worksheet.cell(total_expense_row, 2, BUDGET_TOTAL_LABELS[1])
    worksheet.cell(period_balance_row, 2, BUDGET_TOTAL_LABELS[2])

    for column in range(3, 3 + BUDGET_MONTHS):
        letter = get_column_letter(column)
        income_formula = (
            f"SUM({letter}{income_rows[0]}:{letter}{income_rows[-1]})"
            if income_rows
            else "0"
        )
        expense_formula = (
            f"SUM({letter}{expense_rows[0]}:{letter}{expense_rows[-1]})"
            if expense_rows
            else "0"
        )
        worksheet.cell(total_income_row, column, f"={income_formula}")
        worksheet.cell(total_expense_row, column, f"={expense_formula}")
        worksheet.cell(
            period_balance_row,
            column,
            f"={letter}{total_income_row}+{letter}{total_expense_row}",
        )

    dark_blue = "465B75"
    light_blue = "EAF2F8"
    light_green = "EAF8F1"
    light_red = "FDEDEE"
    thin_gray = Side(style="thin", color="D8DEE7")
    total_border = Border(top=Side(style="thin", color="8795A8"))
    financial_format = '#,##0.00 "M";[Red](#,##0.00 "M");-'

    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.protection = Protection(locked=True)
        cell.border = Border(bottom=thin_gray)

    for row in range(first_data_row, last_data_row + 1):
        worksheet.cell(row, 1).protection = Protection(locked=True)
        worksheet.cell(row, 2).protection = Protection(locked=True)
        for column in range(3, 3 + BUDGET_MONTHS):
            cell = worksheet.cell(row, column)
            cell.protection = Protection(locked=False)
            cell.font = Font(color="0000FF")
            cell.number_format = financial_format
            cell.alignment = Alignment(horizontal="right")

    total_styles = (
        (total_income_row, light_green),
        (total_expense_row, light_red),
        (period_balance_row, light_blue),
    )
    for row, fill_color in total_styles:
        for column in range(2, 3 + BUDGET_MONTHS):
            cell = worksheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(bold=True, color="000000")
            cell.protection = Protection(locked=True)
            cell.border = total_border
            if column >= 3:
                cell.number_format = financial_format

    worksheet.column_dimensions["A"].width = 10
    worksheet.column_dimensions["B"].width = 38
    for column in range(3, 3 + BUDGET_MONTHS):
        worksheet.column_dimensions[get_column_letter(column)].width = 14
    worksheet.row_dimensions[1].height = 24
    worksheet.freeze_panes = "C2"
    if accounts:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(2 + BUDGET_MONTHS)}{last_data_row}"

    worksheet.protection.sheet = True
    worksheet.protection.autoFilter = False
    worksheet.protection.formatColumns = False
    worksheet.protection.selectLockedCells = False
    worksheet.protection.selectUnlockedCells = False
    worksheet.protection.enable()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook


def cash_budget_workbook_bytes(session: Session, year: int) -> bytes:
    output = BytesIO()
    build_cash_budget_workbook(session, year).save(output)
    return output.getvalue()


def _header_month(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date().replace(day=1)
    if isinstance(value, date):
        return value.replace(day=1)
    text = str(value or "").strip()
    for pattern in ("%Y-%m", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date().replace(day=1)
        except ValueError:
            continue
    return None


def _decimal_value(value: object, *, row: int, period: date) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, bool) or (isinstance(value, str) and value.lstrip().startswith("=")):
        raise ValueError(f"Fila {row}, {period:%Y-%m}: el importe debe ser numérico")
    if isinstance(value, float) and not isfinite(value):
        raise ValueError(f"Fila {row}, {period:%Y-%m}: el importe debe ser finito")
    text = str(value).strip().replace("$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        parsed = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Fila {row}, {period:%Y-%m}: el importe debe ser numérico") from exc
    if not parsed.is_finite():
        raise ValueError(f"Fila {row}, {period:%Y-%m}: el importe debe ser finito")
    return (parsed * MILLION).quantize(Decimal("0.01"))


def import_cash_budget_workbook(
    session: Session,
    content: bytes,
    year: int,
) -> dict[str, int | str]:
    if not content:
        raise ValueError("El archivo está vacío")
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    except Exception as exc:
        raise ValueError("El archivo no es un Excel válido") from exc

    if BUDGET_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"No se encontró la hoja '{BUDGET_SHEET_NAME}'")
    worksheet = workbook[BUDGET_SHEET_NAME]
    months = budget_months(year)
    if str(worksheet.cell(1, 1).value or "").strip() != "ID":
        raise ValueError("La primera columna debe ser ID")
    if str(worksheet.cell(1, 2).value or "").strip() != "Cuenta Cash":
        raise ValueError("La segunda columna debe ser Cuenta Cash")
    for offset, expected_month in enumerate(months, start=3):
        if _header_month(worksheet.cell(1, offset).value) != expected_month:
            raise ValueError(
                f"La columna {get_column_letter(offset)} debe corresponder a {expected_month:%Y-%m}"
            )

    accounts = _cash_accounts(session)
    accounts_by_id = {int(account.id): account for account in accounts if account.id is not None}
    imported: dict[tuple[int, date], Decimal] = {}
    seen_ids: set[int] = set()
    for row_number in range(2, worksheet.max_row + 1):
        raw_id = worksheet.cell(row_number, 1).value
        raw_name = worksheet.cell(row_number, 2).value
        name = str(raw_name or "").strip()
        if name in BUDGET_TOTAL_LABELS or (raw_id in (None, "") and not name):
            continue
        try:
            account_id = int(raw_id)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Fila {row_number}: ID de cuenta inválido") from exc
        if account_id in seen_ids:
            raise ValueError(f"Fila {row_number}: la cuenta {account_id} está duplicada")
        account = accounts_by_id.get(account_id)
        if account is None:
            raise ValueError(f"Fila {row_number}: la cuenta Cash {account_id} no existe o no es presupuestable")
        if name != str(account.descripcion).strip():
            raise ValueError(f"Fila {row_number}: el nombre de la cuenta Cash {account_id} fue modificado")
        seen_ids.add(account_id)
        for month_offset, month in enumerate(months, start=3):
            imported[(account_id, month)] = _decimal_value(
                worksheet.cell(row_number, month_offset).value,
                row=row_number,
                period=month,
            )

    missing_ids = set(accounts_by_id) - seen_ids
    if missing_ids:
        raise ValueError(
            "Faltan cuentas Cash de la plantilla: " + ", ".join(str(value) for value in sorted(missing_ids))
        )

    existing_rows = session.exec(
        select(ErpCashProyectado).where(
            ErpCashProyectado.deleted_at.is_(None),
            ErpCashProyectado.tipo == BUDGET_TYPE,
            ErpCashProyectado.cuenta_cash_id.in_(list(accounts_by_id)),
            ErpCashProyectado.fecha_periodo >= months[0],
            ErpCashProyectado.fecha_periodo <= months[-1],
        )
    ).all()
    existing_by_key = {
        (int(row.cuenta_cash_id), row.fecha_periodo): row for row in existing_rows
    }
    created = updated = unchanged = 0
    now = current_utc_time()
    try:
        for key, amount in imported.items():
            existing = existing_by_key.get(key)
            if existing is None:
                session.add(
                    ErpCashProyectado(
                        cuenta_cash_id=key[0],
                        fecha_periodo=key[1],
                        tipo=BUDGET_TYPE,
                        importe=amount,
                        observacion="IMPORTADO DESDE PRESUPUESTO XLSX",
                    )
                )
                created += 1
            elif Decimal(existing.importe or 0) != amount:
                existing.importe = amount
                existing.observacion = "IMPORTADO DESDE PRESUPUESTO XLSX"
                existing.updated_at = now
                existing.version = int(existing.version or 0) + 1
                session.add(existing)
                updated += 1
            else:
                unchanged += 1
        session.commit()
    except Exception:
        session.rollback()
        raise

    return {
        "year": year,
        "months": BUDGET_MONTHS,
        "accounts": len(accounts_by_id),
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
    }
