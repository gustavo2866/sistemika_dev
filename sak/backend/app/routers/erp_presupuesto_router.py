from collections import defaultdict
from io import BytesIO
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Literal
from urllib.parse import quote

import psycopg
from fastapi import Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from pydantic import BaseModel, Field
from sqlalchemy import String, and_, cast, extract, func
from sqlmodel import Session, select

from app.core.router import create_generic_router
from app.core.nested_crud import NestedCRUD
from app.db import get_session
from app.models.constructora.proyectos_conceptos import ProyectosConceptos
from app.models.erp.cuenta import ErpCuenta
from app.models.erp.libro_diario import ErpLibroDiario
from app.models.erp.libro_diario_sync import get_dest_url, sync_real_presupuestos
from app.models.erp.presupuesto import ErpPresupuesto
from app.models.erp.rubro import ErpRubro
from app.models.proyecto import Proyecto

def _find_same_month_presupuesto(
    session: Session,
    *,
    proyecto_id: int,
    erp_cuenta_id: int,
    fecha: date,
    exclude_id: int | None = None,
) -> ErpPresupuesto | None:
    stmt = (
        select(ErpPresupuesto)
        .where(ErpPresupuesto.deleted_at.is_(None))
        .where(ErpPresupuesto.proyecto_id == proyecto_id)
        .where(ErpPresupuesto.erp_cuenta_id == erp_cuenta_id)
        .where(extract("year", ErpPresupuesto.fecha) == fecha.year)
        .where(extract("month", ErpPresupuesto.fecha) == fecha.month)
    )
    if exclude_id is not None:
        stmt = stmt.where(ErpPresupuesto.id != exclude_id)
    return session.exec(stmt).first()


def _validate_unique_presupuesto_month(
    session: Session,
    *,
    proyecto_id: int,
    erp_cuenta_id: int,
    fecha: date,
    exclude_id: int | None = None,
) -> None:
    duplicate = _find_same_month_presupuesto(
        session,
        proyecto_id=proyecto_id,
        erp_cuenta_id=erp_cuenta_id,
        fecha=fecha,
        exclude_id=exclude_id,
    )
    if duplicate is None:
        return
    raise HTTPException(
        status_code=409,
        detail={
            "error": {
                "code": "DUPLICATE_ERP_PRESUPUESTO_MONTH",
                "message": (
                    "Ya existe un presupuesto para este proyecto, cuenta ERP y mes."
                ),
                "details": {
                    "existing_id": duplicate.id,
                    "proyecto_id": proyecto_id,
                    "erp_cuenta_id": erp_cuenta_id,
                    "periodo": f"{fecha.year:04d}-{fecha.month:02d}",
                },
            }
        },
    )


class ErpPresupuestoCRUD(NestedCRUD):
    def create(self, session: Session, data: dict):
        fecha_value = data.get("fecha")
        fecha = date.fromisoformat(fecha_value) if isinstance(fecha_value, str) else fecha_value
        if isinstance(fecha, date):
            _validate_unique_presupuesto_month(
                session,
                proyecto_id=int(data["proyecto_id"]),
                erp_cuenta_id=int(data["erp_cuenta_id"]),
                fecha=fecha,
            )
        return super().create(session, data)

    def update(
        self,
        session: Session,
        obj_id: int,
        data: dict,
        check_version: bool = True,
    ):
        current = self.get(session, obj_id)
        if current is None:
            return None

        fecha_value = data.get("fecha", current.fecha)
        fecha = date.fromisoformat(fecha_value) if isinstance(fecha_value, str) else fecha_value
        proyecto_id = int(data.get("proyecto_id", current.proyecto_id))
        erp_cuenta_id = int(data.get("erp_cuenta_id", current.erp_cuenta_id))
        if isinstance(fecha, date):
            _validate_unique_presupuesto_month(
                session,
                proyecto_id=proyecto_id,
                erp_cuenta_id=erp_cuenta_id,
                fecha=fecha,
                exclude_id=int(obj_id),
            )
        return super().update(session, obj_id, data, check_version=check_version)


erp_presupuesto_crud = ErpPresupuestoCRUD(ErpPresupuesto, nested_relations={})

erp_presupuesto_router = create_generic_router(
    model=ErpPresupuesto,
    crud=erp_presupuesto_crud,
    prefix="/erp/presupuestos",
    tags=["erp-presupuestos"],
)


class PresupuestoCopyVariation(BaseModel):
    concepto_id: int
    porcentaje: Decimal = Decimal("0")


class PresupuestoCopyRequest(BaseModel):
    proyecto_id: int
    periodo_origen: str
    periodo_destino: str
    variaciones: list[PresupuestoCopyVariation] = Field(default_factory=list)


class PresupuestoClearRequest(BaseModel):
    proyecto_id: int
    periodo: str


class PresupuestoRealIncomeRequest(BaseModel):
    proyecto_id: int
    periodo: str
    erp_cuenta_id: int
    real_ingreso: Decimal = Field(ge=Decimal("0"))


def _decimal_to_float(value: Decimal | int | float | None) -> float:
    if value is None:
        return 0.0
    return float(value)


def _month_key(year_value: object, month_value: object) -> str:
    year = int(year_value)
    month = int(month_value)
    return f"{year:04d}-{month:02d}"


def _parse_month_key(month_key: str) -> tuple[int, int]:
    try:
        year_text, month_text = month_key.split("-", maxsplit=1)
        year = int(year_text)
        month = int(month_text)
    except ValueError as exc:
        raise ValueError("periodo debe tener formato YYYY-MM") from exc

    if month < 1 or month > 12:
        raise ValueError("periodo debe tener un mes entre 01 y 12")

    return year, month


def _apply_percent_variation(value: Decimal | int | float | None, percent: Decimal) -> Decimal:
    base = Decimal("0") if value is None else Decimal(str(value))
    multiplier = Decimal("1") + (percent / Decimal("100"))
    return (base * multiplier).quantize(Decimal("0.01"))


def _build_panel_movimientos_stmt(
    anio: int,
    mes: int,
    proyecto_id: int,
    concepto: Literal["egreso", "ingreso", "resultado"],
    rubro_id: int | None = None,
    erp_cuenta_id: int | None = None,
):
    rubro_nombre_expr = func.lower(func.coalesce(ErpRubro.nombre, ""))

    stmt = (
        select(
            ErpLibroDiario.id,
            ErpLibroDiario.fecha,
            ErpLibroDiario.tipo_asiento,
            ErpLibroDiario.nro_asiento,
            ErpLibroDiario.descripcion,
            ErpLibroDiario.debe,
            ErpLibroDiario.haber,
            ErpCuenta.nro_cuenta,
            ErpCuenta.cod_cuenta,
            ErpCuenta.descripcion.label("cuenta_descripcion"),
        )
        .join(
            Proyecto,
            (Proyecto.deleted_at.is_(None))
            & (Proyecto.centro_costo.is_not(None))
            & (cast(Proyecto.centro_costo, String) == func.btrim(ErpLibroDiario.centro_costo)),
        )
        .join(
            ErpCuenta,
            (ErpCuenta.deleted_at.is_(None))
            & (ErpCuenta.nro_cuenta == ErpLibroDiario.cuenta_codigo),
        )
        .join(ErpRubro, ErpRubro.id == ErpCuenta.rubro_id)
        .where(ErpLibroDiario.deleted_at.is_(None))
        .where(ErpLibroDiario.periodo_anio == anio)
        .where(ErpLibroDiario.periodo_mes == mes)
        .where(Proyecto.id == proyecto_id)
        .order_by(
            ErpLibroDiario.fecha,
            ErpLibroDiario.tipo_asiento,
            ErpLibroDiario.nro_asiento,
            ErpLibroDiario.nro_renglon,
            ErpLibroDiario.id,
        )
    )

    if rubro_id is not None:
        stmt = stmt.where(ErpRubro.id == rubro_id)

    if erp_cuenta_id is not None:
        stmt = stmt.where(ErpCuenta.id == erp_cuenta_id)

    if concepto == "ingreso":
        stmt = stmt.where(rubro_nombre_expr == "ingresos")
    elif concepto == "egreso":
        stmt = stmt.where(rubro_nombre_expr != "ingresos")

    return stmt


def _format_movimiento_asiento(row) -> str:
    return "-".join(str(value) for value in [row.tipo_asiento, row.nro_asiento] if value) or ""


def _format_movimiento_cuenta(row) -> str:
    return " - ".join(
        str(value)
        for value in [row.nro_cuenta, row.cod_cuenta, row.cuenta_descripcion]
        if value
    )


def _empty_month_values(months: list[str]) -> dict[str, dict]:
    return {
        month: {
            "ingresos": 0.0,
            "egresos": 0.0,
            "real_ingresos": 0.0,
            "real_egresos": 0.0,
            "empleados": 0.0,
            "presupuesto_id": None,
            "record_count": 0,
        }
        for month in months
    }


def _add_month_values(
    target: dict[str, dict],
    month: str,
    ingresos: Decimal | int | float | None,
    egresos: Decimal | int | float | None,
    real_ingresos: Decimal | int | float | None,
    real_egresos: Decimal | int | float | None,
    empleados: Decimal | int | float | None,
    presupuesto_id: int | None,
    record_count: int,
) -> None:
    if month not in target:
        target[month] = {
            "ingresos": 0.0,
            "egresos": 0.0,
            "real_ingresos": 0.0,
            "real_egresos": 0.0,
            "empleados": 0.0,
            "presupuesto_id": None,
            "record_count": 0,
        }

    target[month]["ingresos"] += _decimal_to_float(ingresos)
    target[month]["egresos"] += _decimal_to_float(egresos)
    target[month]["real_ingresos"] += _decimal_to_float(real_ingresos)
    target[month]["real_egresos"] += _decimal_to_float(real_egresos)
    target[month]["empleados"] += _decimal_to_float(empleados)
    target[month]["record_count"] += record_count

    if target[month]["record_count"] == 1 and presupuesto_id is not None:
        target[month]["presupuesto_id"] = presupuesto_id
    elif target[month]["presupuesto_id"] != presupuesto_id:
        target[month]["presupuesto_id"] = None


@erp_presupuesto_router.get("/panel")
def get_erp_presupuesto_panel(
    fecha_desde: date = Query(..., description="Primer dia del periodo"),
    fecha_hasta: date = Query(..., description="Ultimo dia del periodo"),
    estado: str = Query("02-ejecucion", description="Estado de proyecto a incluir"),
    proyecto_id: int | None = Query(None, description="Filtro opcional por proyecto"),
    session: Session = Depends(get_session),
):
    """Devuelve una vista agregada para el panel de presupuestos ERP."""
    year_expr = extract("year", ErpPresupuesto.fecha).label("year")
    month_expr = extract("month", ErpPresupuesto.fecha).label("month")
    ingresos_expr = func.coalesce(func.sum(ErpPresupuesto.ingres), 0).label("ingresos")
    egresos_expr = func.coalesce(func.sum(ErpPresupuesto.egreso), 0).label("egresos")
    real_ingresos_expr = func.coalesce(func.sum(ErpPresupuesto.real_ingreso), 0).label("real_ingresos")
    real_egresos_expr = func.coalesce(func.sum(ErpPresupuesto.real_egreso), 0).label("real_egresos")
    empleados_expr = func.coalesce(func.sum(ErpPresupuesto.obreros_cantidad), 0).label("empleados")
    presupuesto_id_expr = func.min(ErpPresupuesto.id).label("presupuesto_id")
    record_count_expr = func.count(ErpPresupuesto.id).label("record_count")

    stmt = (
        select(
            Proyecto.id.label("proyecto_id"),
            Proyecto.nombre.label("proyecto_nombre"),
            ErpRubro.id.label("rubro_id"),
            ErpRubro.nombre.label("rubro_nombre"),
            ErpCuenta.id.label("cuenta_id"),
            ErpCuenta.cod_cuenta.label("cuenta_codigo"),
            ErpCuenta.descripcion.label("cuenta_nombre"),
            ErpCuenta.nro_cuenta.label("cuenta_orden"),
            year_expr,
            month_expr,
            ingresos_expr,
            egresos_expr,
            real_ingresos_expr,
            real_egresos_expr,
            empleados_expr,
            presupuesto_id_expr,
            record_count_expr,
        )
        .join(Proyecto, ErpPresupuesto.proyecto_id == Proyecto.id)
        .join(ErpCuenta, ErpPresupuesto.erp_cuenta_id == ErpCuenta.id)
        .join(ErpRubro, ErpCuenta.rubro_id == ErpRubro.id)
        .where(ErpPresupuesto.fecha >= fecha_desde)
        .where(ErpPresupuesto.fecha <= fecha_hasta)
        .where(Proyecto.estado == estado)
        .where(ErpPresupuesto.deleted_at.is_(None))
        .group_by(
            Proyecto.id,
            Proyecto.nombre,
            ErpRubro.id,
            ErpRubro.nombre,
            ErpCuenta.id,
            ErpCuenta.cod_cuenta,
            ErpCuenta.descripcion,
            ErpCuenta.nro_cuenta,
            year_expr,
            month_expr,
        )
        .order_by(Proyecto.nombre, ErpRubro.nombre, ErpCuenta.nro_cuenta, ErpCuenta.cod_cuenta)
    )

    if proyecto_id is not None:
        stmt = stmt.where(Proyecto.id == proyecto_id)

    rows = session.exec(stmt).all()

    months: list[str] = []
    cursor = date(fecha_desde.year, fecha_desde.month, 1)
    end_cursor = date(fecha_hasta.year, fecha_hasta.month, 1)
    while cursor <= end_cursor:
        months.append(f"{cursor.year:04d}-{cursor.month:02d}")
        if cursor.month == 12:
            cursor = date(cursor.year + 1, 1, 1)
        else:
            cursor = date(cursor.year, cursor.month + 1, 1)

    projects: dict[int, dict] = {}
    rubros_by_project: dict[int, dict[int, dict]] = defaultdict(dict)
    cuentas_by_rubro: dict[tuple[int, int], dict[int, dict]] = defaultdict(dict)

    for row in rows:
        month = _month_key(row.year, row.month)
        project = projects.get(row.proyecto_id)
        if project is None:
            project = {
                "proyecto_id": row.proyecto_id,
                "proyecto_nombre": row.proyecto_nombre,
                "months": _empty_month_values(months),
                "rubros": [],
            }
            projects[row.proyecto_id] = project

        rubro = rubros_by_project[row.proyecto_id].get(row.rubro_id)
        if rubro is None:
            rubro = {
                "rubro_id": row.rubro_id,
                "rubro_nombre": row.rubro_nombre,
                "months": _empty_month_values(months),
                "cuentas": [],
            }
            rubros_by_project[row.proyecto_id][row.rubro_id] = rubro
            project["rubros"].append(rubro)

        cuenta_key = (row.proyecto_id, row.rubro_id)
        cuenta = cuentas_by_rubro[cuenta_key].get(row.cuenta_id)
        if cuenta is None:
            cuenta = {
                "cuenta_id": row.cuenta_id,
                "cuenta_codigo": row.cuenta_codigo,
                "cuenta_nombre": row.cuenta_nombre,
                "months": _empty_month_values(months),
            }
            cuentas_by_rubro[cuenta_key][row.cuenta_id] = cuenta
            rubro["cuentas"].append(cuenta)

        record_count = int(row.record_count or 0)
        project_real_ingresos = (
            0
            if str(row.rubro_nombre or "").strip().lower() == "ingresos"
            else row.real_ingresos
        )
        _add_month_values(
            project["months"],
            month,
            row.ingresos,
            row.egresos,
            project_real_ingresos,
            row.real_egresos,
            row.empleados,
            row.presupuesto_id,
            record_count,
        )
        _add_month_values(
            rubro["months"],
            month,
            row.ingresos,
            row.egresos,
            row.real_ingresos,
            row.real_egresos,
            row.empleados,
            row.presupuesto_id,
            record_count,
        )
        _add_month_values(
            cuenta["months"],
            month,
            row.ingresos,
            row.egresos,
            row.real_ingresos,
            row.real_egresos,
            row.empleados,
            row.presupuesto_id,
            record_count,
        )

    return {
        "fecha_desde": fecha_desde.isoformat(),
        "fecha_hasta": fecha_hasta.isoformat(),
        "estado": estado,
        "months": months,
        "rows": list(projects.values()),
    }


@erp_presupuesto_router.get("/panel/proyectos")
def get_erp_presupuesto_panel_proyectos(
    estado: str = Query("02-ejecucion", description="Estado de proyecto a incluir"),
    session: Session = Depends(get_session),
):
    """Devuelve proyectos disponibles para filtrar el panel de presupuestos ERP."""
    rows = session.exec(
        select(Proyecto.id, Proyecto.nombre)
        .where(Proyecto.deleted_at.is_(None))
        .where(Proyecto.estado == estado)
        .order_by(Proyecto.nombre)
    ).all()

    return {
        "estado": estado,
        "rows": [
            {
                "proyecto_id": row.id,
                "proyecto_nombre": row.nombre,
            }
            for row in rows
        ],
    }


@erp_presupuesto_router.get("/panel/export")
def export_erp_presupuesto_panel_project(
    periodo: str = Query(..., description="Periodo en formato YYYY-MM"),
    proyecto_id: int = Query(..., description="Proyecto a exportar"),
    session: Session = Depends(get_session),
):
    """Exporta el presupuesto de un proyecto y periodo a Excel."""
    try:
        anio, mes = _parse_month_key(periodo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    proyecto = session.get(Proyecto, proyecto_id)
    if proyecto is None or proyecto.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")

    rows = session.exec(
        select(
            ErpRubro.id.label("rubro_id"),
            ErpRubro.nombre.label("rubro_nombre"),
            ErpCuenta.cod_cuenta.label("cuenta_codigo"),
            ErpCuenta.descripcion.label("cuenta_nombre"),
            ErpCuenta.nro_cuenta.label("cuenta_orden"),
            func.coalesce(func.sum(ErpPresupuesto.ingres), 0).label("ingres"),
            func.coalesce(func.sum(ErpPresupuesto.real_ingreso), 0).label("real_ingresos"),
            func.coalesce(func.sum(ErpPresupuesto.egreso), 0).label("egreso"),
            func.coalesce(func.sum(ErpPresupuesto.obreros_cantidad), 0).label("obreros_cantidad"),
        )
        .select_from(ErpCuenta)
        .join(ErpRubro, ErpRubro.id == ErpCuenta.rubro_id)
        .outerjoin(
            ErpPresupuesto,
            and_(
                ErpPresupuesto.erp_cuenta_id == ErpCuenta.id,
                ErpPresupuesto.proyecto_id == proyecto_id,
                ErpPresupuesto.deleted_at.is_(None),
                extract("year", ErpPresupuesto.fecha) == anio,
                extract("month", ErpPresupuesto.fecha) == mes,
            ),
        )
        .where(ErpCuenta.deleted_at.is_(None))
        .where(ErpRubro.deleted_at.is_(None))
        .group_by(
            ErpRubro.id,
            ErpRubro.nombre,
            ErpCuenta.cod_cuenta,
            ErpCuenta.descripcion,
            ErpCuenta.nro_cuenta,
        )
        .order_by(ErpCuenta.cod_cuenta, ErpCuenta.descripcion, ErpRubro.nombre)
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    headers = [
        "centro costo",
        "periodo",
        "rubro",
        "cuenta",
        "empleados",
        "egresos",
        "ingresos",
        "real_ingresos",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="E5E7EB")
    header_font = Font(bold=True, color="111827")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.protection = Protection(locked=True)

    for row in rows:
        rubro_nombre = row.rubro_nombre or ""
        rubro_label = f"{row.rubro_id} - {rubro_nombre}"
        proyecto_nombre = proyecto.nombre or ""
        centro_costo_label = f"{proyecto.centro_costo} - {proyecto_nombre[:30]}"
        cuenta_label = " - ".join(
            value
            for value in [row.cuenta_codigo, row.cuenta_nombre]
            if value
        )
        ws.append(
            [
                centro_costo_label,
                periodo,
                rubro_label,
                cuenta_label,
                float(row.obreros_cantidad or 0),
                float(row.egreso or 0),
                float(row.ingres or 0),
                float(row.real_ingresos or 0),
            ]
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.protection = Protection(locked=False)
        row[0].protection = Protection(locked=True)
        row[2].protection = Protection(locked=True)
        row[3].protection = Protection(locked=True)
        row[4].number_format = '#,##0.00'
        row[5].number_format = '$ #,##0.00'
        row[6].number_format = '$ #,##0.00'
        row[7].number_format = '$ #,##0.00'

    widths = {
        "A": 42,
        "B": 12,
        "C": 34,
        "D": 58,
        "E": 12,
        "F": 16,
        "G": 16,
        "H": 16,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.enable()

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"presupuesto-{proyecto_id}-{periodo}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
        },
    )


@erp_presupuesto_router.post("/panel/copy")
def copy_erp_presupuesto_panel_project(
    payload: PresupuestoCopyRequest,
    session: Session = Depends(get_session),
):
    """Copia el presupuesto de un periodo origen al periodo destino del mismo proyecto."""
    try:
        origen_anio, origen_mes = _parse_month_key(payload.periodo_origen)
        destino_anio, destino_mes = _parse_month_key(payload.periodo_destino)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    if payload.periodo_origen == payload.periodo_destino:
        raise HTTPException(status_code=400, detail="El periodo origen y destino deben ser distintos.")

    proyecto = session.get(Proyecto, payload.proyecto_id)
    if proyecto is None or proyecto.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")

    variaciones_by_concepto: dict[int, Decimal] = {}
    for variacion in payload.variaciones:
        porcentaje = variacion.porcentaje
        if not porcentaje.is_finite():
            raise HTTPException(status_code=400, detail="El porcentaje de variacion debe ser finito.")
        if porcentaje < Decimal("-100"):
            raise HTTPException(status_code=400, detail="El porcentaje de variacion no puede ser menor a -100.")
        variaciones_by_concepto[int(variacion.concepto_id)] = porcentaje

    if variaciones_by_concepto:
        existing_concepts = session.exec(
            select(ProyectosConceptos.id).where(
                ProyectosConceptos.id.in_(list(variaciones_by_concepto.keys()))
            )
        ).all()
        missing_concepts = set(variaciones_by_concepto) - {int(concept_id) for concept_id in existing_concepts}
        if missing_concepts:
            raise HTTPException(status_code=400, detail="Uno o mas conceptos no existen.")

    source_rows = session.exec(
        select(
            ErpPresupuesto.erp_cuenta_id.label("erp_cuenta_id"),
            ErpCuenta.proyectos_concepto_id.label("proyectos_concepto_id"),
            func.coalesce(func.sum(ErpPresupuesto.egreso), 0).label("egreso"),
            func.coalesce(func.sum(ErpPresupuesto.ingres), 0).label("ingres"),
            func.coalesce(func.sum(ErpPresupuesto.obreros_cantidad), 0).label("obreros_cantidad"),
            func.coalesce(func.sum(ErpPresupuesto.obreros_costo), 0).label("obreros_costo"),
        )
        .join(ErpCuenta, ErpPresupuesto.erp_cuenta_id == ErpCuenta.id)
        .where(ErpPresupuesto.deleted_at.is_(None))
        .where(ErpCuenta.deleted_at.is_(None))
        .where(ErpPresupuesto.proyecto_id == payload.proyecto_id)
        .where(extract("year", ErpPresupuesto.fecha) == origen_anio)
        .where(extract("month", ErpPresupuesto.fecha) == origen_mes)
        .group_by(ErpPresupuesto.erp_cuenta_id, ErpCuenta.proyectos_concepto_id)
    ).all()

    if not source_rows:
        raise HTTPException(status_code=404, detail="No hay presupuesto en el periodo origen.")

    source_rows = [
        row
        for row in source_rows
        if any(
            Decimal(str(value or 0)) != 0
            for value in [row.egreso, row.ingres, row.obreros_cantidad, row.obreros_costo]
        )
    ]
    if not source_rows:
        raise HTTPException(status_code=404, detail="El periodo origen no tiene importes presupuestados.")

    existing_rows = session.exec(
        select(ErpPresupuesto)
        .where(ErpPresupuesto.deleted_at.is_(None))
        .where(ErpPresupuesto.proyecto_id == payload.proyecto_id)
        .where(extract("year", ErpPresupuesto.fecha) == destino_anio)
        .where(extract("month", ErpPresupuesto.fecha) == destino_mes)
    ).all()
    for presupuesto in existing_rows:
        session.delete(presupuesto)

    destino_fecha = date(destino_anio, destino_mes, 1)
    copied_rows = 0
    varied_rows = 0
    for row in source_rows:
        concepto_id = row.proyectos_concepto_id
        porcentaje = (
            variaciones_by_concepto.get(int(concepto_id), Decimal("0"))
            if concepto_id is not None
            else Decimal("0")
        )
        egreso = _apply_percent_variation(row.egreso, porcentaje)
        ingres = _apply_percent_variation(row.ingres, porcentaje)
        obreros_cantidad = _apply_percent_variation(row.obreros_cantidad, porcentaje)
        obreros_costo = _apply_percent_variation(row.obreros_costo, porcentaje)

        if egreso == 0 and ingres == 0 and obreros_cantidad == 0 and obreros_costo == 0:
            continue

        if porcentaje != 0:
            varied_rows += 1
        copied_rows += 1
        session.add(
            ErpPresupuesto(
                fecha=destino_fecha,
                proyecto_id=payload.proyecto_id,
                erp_cuenta_id=int(row.erp_cuenta_id),
                egreso=egreso,
                ingres=ingres,
                obreros_cantidad=obreros_cantidad,
                obreros_costo=obreros_costo,
            )
        )

    session.commit()

    with psycopg.connect(get_dest_url()) as conn:
        sync_result = sync_real_presupuestos(conn, destino_anio, destino_mes, dry_run=False)
        conn.commit()

    return {
        "periodo_origen": payload.periodo_origen,
        "periodo_destino": payload.periodo_destino,
        "proyecto_id": payload.proyecto_id,
        "proyecto_nombre": proyecto.nombre,
        "source_rows": len(source_rows),
        "copied_rows": copied_rows,
        "deleted_rows": len(existing_rows),
        "varied_rows": varied_rows,
        "sync": sync_result,
    }


@erp_presupuesto_router.post("/panel/clear")
def clear_erp_presupuesto_panel_project(
    payload: PresupuestoClearRequest,
    session: Session = Depends(get_session),
):
    """Limpia el presupuesto de un proyecto para un periodo."""
    try:
        anio, mes = _parse_month_key(payload.periodo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    proyecto = session.get(Proyecto, payload.proyecto_id)
    if proyecto is None or proyecto.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")

    existing_rows = session.exec(
        select(ErpPresupuesto)
        .where(ErpPresupuesto.deleted_at.is_(None))
        .where(ErpPresupuesto.proyecto_id == payload.proyecto_id)
        .where(extract("year", ErpPresupuesto.fecha) == anio)
        .where(extract("month", ErpPresupuesto.fecha) == mes)
    ).all()

    for presupuesto in existing_rows:
        session.delete(presupuesto)

    session.commit()

    with psycopg.connect(get_dest_url()) as conn:
        sync_result = sync_real_presupuestos(conn, anio, mes, dry_run=False)
        conn.commit()

    return {
        "periodo": payload.periodo,
        "proyecto_id": payload.proyecto_id,
        "proyecto_nombre": proyecto.nombre,
        "deleted_rows": len(existing_rows),
        "sync": sync_result,
    }


@erp_presupuesto_router.post("/panel/real-income")
def set_erp_presupuesto_panel_real_income(
    payload: PresupuestoRealIncomeRequest,
    session: Session = Depends(get_session),
):
    """Setea el ingreso real de una cuenta/proyecto/periodo desde el panel."""
    try:
        anio, mes = _parse_month_key(payload.periodo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    proyecto = session.get(Proyecto, payload.proyecto_id)
    if proyecto is None or proyecto.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")

    cuenta = session.get(ErpCuenta, payload.erp_cuenta_id)
    if cuenta is None or cuenta.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Cuenta ERP no encontrada")

    rows = session.exec(
        select(ErpPresupuesto)
        .where(ErpPresupuesto.deleted_at.is_(None))
        .where(ErpPresupuesto.proyecto_id == payload.proyecto_id)
        .where(ErpPresupuesto.erp_cuenta_id == payload.erp_cuenta_id)
        .where(extract("year", ErpPresupuesto.fecha) == anio)
        .where(extract("month", ErpPresupuesto.fecha) == mes)
        .order_by(ErpPresupuesto.fecha, ErpPresupuesto.id)
    ).all()

    updated_at = datetime.now(UTC)
    if len(rows) > 1:
        raise HTTPException(
            status_code=409,
            detail={
                "error": {
                    "code": "DUPLICATE_ERP_PRESUPUESTO_MONTH",
                    "message": (
                        "Hay mas de un presupuesto para este proyecto, cuenta ERP y mes."
                    ),
                    "details": {
                        "ids": [row.id for row in rows],
                        "proyecto_id": payload.proyecto_id,
                        "erp_cuenta_id": payload.erp_cuenta_id,
                        "periodo": payload.periodo,
                    },
                }
            },
        )

    if rows:
        target = rows[0]
        target.real_ingreso = payload.real_ingreso
        if hasattr(target, "updated_at"):
            target.updated_at = updated_at
        if hasattr(target, "version"):
            target.version += 1
        session.add(target)
    else:
        target = ErpPresupuesto(
            fecha=date(anio, mes, 1),
            proyecto_id=payload.proyecto_id,
            erp_cuenta_id=payload.erp_cuenta_id,
            real_ingreso=payload.real_ingreso,
        )
        session.add(target)

    session.commit()
    session.refresh(target)

    return {
        "periodo": payload.periodo,
        "proyecto_id": payload.proyecto_id,
        "erp_cuenta_id": payload.erp_cuenta_id,
        "presupuesto_id": target.id,
        "real_ingreso": _decimal_to_float(target.real_ingreso),
        "affected_rows": max(len(rows), 1),
    }


EXPECTED_PRESUPUESTO_IMPORT_HEADERS = [
    "centro costo",
    "periodo",
    "rubro",
    "cuenta",
    "empleados",
    "egresos",
    "ingresos",
    "real_ingresos",
]


def _parse_excel_decimal(value: object, row_number: int, column_name: str, errors: list[str]) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    try:
        return Decimal(str(value).replace("$", "").replace(".", "").replace(",", ".").strip())
    except Exception:
        errors.append(f"Fila {row_number}: {column_name} debe ser numerico.")
        return Decimal("0")


@erp_presupuesto_router.post("/panel/import")
async def import_erp_presupuesto_panel_project(
    periodo: str = Query(..., description="Periodo en formato YYYY-MM"),
    proyecto_id: int = Query(..., description="Proyecto a importar"),
    process: bool = Form(False, description="Si es true, procesa luego de validar"),
    file: UploadFile = File(...),
    session: Session = Depends(get_session),
):
    """Valida o importa un presupuesto Excel exportado desde el panel."""
    try:
        anio, mes = _parse_month_key(periodo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    proyecto = session.get(Proyecto, proyecto_id)
    if proyecto is None or proyecto.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="El archivo no es un Excel valido.") from exc

    ws = wb.active
    errors: list[str] = []
    headers = [
        str(ws.cell(1, col).value or "").strip().lower()
        for col in range(1, len(EXPECTED_PRESUPUESTO_IMPORT_HEADERS) + 1)
    ]
    if headers != EXPECTED_PRESUPUESTO_IMPORT_HEADERS:
        errors.append(
            "La cabecera no coincide. Se esperan columnas: "
            + ", ".join(EXPECTED_PRESUPUESTO_IMPORT_HEADERS)
        )

    cuentas = session.exec(
        select(ErpCuenta.id, ErpCuenta.cod_cuenta)
        .where(ErpCuenta.deleted_at.is_(None))
    ).all()
    cuenta_id_by_codigo = {str(row.cod_cuenta).strip(): row.id for row in cuentas}
    expected_centro_prefix = str(proyecto.centro_costo)
    imported_rows: list[dict[str, Decimal | int]] = []
    seen_cuentas: set[int] = set()

    for row_number in range(2, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(1, 9)]
        if all(value is None or value == "" for value in values):
            continue

        (
            centro_costo,
            row_periodo,
            _rubro,
            cuenta_label,
            empleados_raw,
            egresos_raw,
            ingresos_raw,
            _real_ingresos_raw,
        ) = values
        centro_text = str(centro_costo or "").strip()
        periodo_text = str(row_periodo or "").strip()
        cuenta_text = str(cuenta_label or "").strip()
        cuenta_codigo = cuenta_text.split(" - ", 1)[0].strip()

        if not centro_text.startswith(expected_centro_prefix):
            errors.append(
                f"Fila {row_number}: centro costo '{centro_text}' no corresponde al proyecto seleccionado."
            )
        if periodo_text != periodo:
            errors.append(f"Fila {row_number}: periodo '{periodo_text}' no coincide con {periodo}.")
        cuenta_id = cuenta_id_by_codigo.get(cuenta_codigo)
        if cuenta_id is None:
            errors.append(f"Fila {row_number}: cuenta '{cuenta_codigo}' no existe en erp_cuentas.")
            continue
        if cuenta_id in seen_cuentas:
            errors.append(f"Fila {row_number}: cuenta '{cuenta_codigo}' duplicada.")
        seen_cuentas.add(cuenta_id)

        empleados = _parse_excel_decimal(empleados_raw, row_number, "empleados", errors)
        egresos = _parse_excel_decimal(egresos_raw, row_number, "egresos", errors)
        ingresos = _parse_excel_decimal(ingresos_raw, row_number, "ingresos", errors)
        real_ingresos = _parse_excel_decimal(
            _real_ingresos_raw, row_number, "real_ingresos", errors
        )
        if empleados < 0 or egresos < 0 or ingresos < 0 or real_ingresos < 0:
            errors.append(
                f"Fila {row_number}: empleados, egresos, ingresos y real_ingresos no pueden ser negativos."
            )

        if empleados != 0 or egresos != 0 or ingresos != 0 or real_ingresos != 0:
            imported_rows.append(
                {
                    "erp_cuenta_id": cuenta_id,
                    "empleados": empleados,
                    "egresos": egresos,
                    "ingresos": ingresos,
                    "real_ingresos": real_ingresos,
                }
            )

    result = {
        "valid": len(errors) == 0,
        "errors": errors,
        "rows_read": max(ws.max_row - 1, 0),
        "rows_budgeted": len(imported_rows),
        "processed": False,
    }

    if errors or not process:
        return result

    period_date = date(anio, mes, 1)
    existing_rows = session.exec(
        select(ErpPresupuesto)
        .where(ErpPresupuesto.deleted_at.is_(None))
        .where(ErpPresupuesto.proyecto_id == proyecto_id)
        .where(extract("year", ErpPresupuesto.fecha) == anio)
        .where(extract("month", ErpPresupuesto.fecha) == mes)
    ).all()
    for presupuesto in existing_rows:
        session.delete(presupuesto)

    for row in imported_rows:
        session.add(
            ErpPresupuesto(
                fecha=period_date,
                proyecto_id=proyecto_id,
                erp_cuenta_id=int(row["erp_cuenta_id"]),
                egreso=row["egresos"],
                ingres=row["ingresos"],
                real_ingreso=row["real_ingresos"],
                obreros_cantidad=row["empleados"],
                obreros_costo=Decimal("0"),
            )
        )
    session.commit()

    with psycopg.connect(get_dest_url()) as conn:
        sync_real_presupuestos(conn, anio, mes, dry_run=False)
        conn.commit()

    result["processed"] = True
    result["deleted_rows"] = len(existing_rows)
    return result


@erp_presupuesto_router.get("/panel/movimientos")
def get_erp_presupuesto_panel_movimientos(
    periodo: str = Query(..., description="Periodo en formato YYYY-MM"),
    proyecto_id: int = Query(..., description="Proyecto a consultar"),
    concepto: Literal["egreso", "ingreso", "resultado"] = Query(
        ..., description="Concepto real a descomponer"
    ),
    rubro_id: int | None = Query(None, description="Filtro opcional por rubro"),
    erp_cuenta_id: int | None = Query(None, description="Filtro opcional por cuenta"),
    session: Session = Depends(get_session),
):
    """Devuelve los movimientos de libro diario que componen un importe real del panel."""
    try:
        anio, mes = _parse_month_key(periodo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    rows = session.exec(
        _build_panel_movimientos_stmt(
            anio,
            mes,
            proyecto_id,
            concepto,
            rubro_id,
            erp_cuenta_id,
        )
    ).all()
    total_debe = sum((row.debe or Decimal("0")) for row in rows)
    total_haber = sum((row.haber or Decimal("0")) for row in rows)

    return {
        "periodo": periodo,
        "concepto": concepto,
        "rows": [
            {
                "id": row.id,
                "fecha": row.fecha.isoformat() if row.fecha else None,
                "tipo_asiento": row.tipo_asiento,
                "nro_asiento": row.nro_asiento,
                "descripcion": row.descripcion,
                "debe": _decimal_to_float(row.debe),
                "haber": _decimal_to_float(row.haber),
                "cuenta": str(row.nro_cuenta) if row.nro_cuenta is not None else "",
                "cuenta_codigo": row.cod_cuenta,
                "cuenta_descripcion": row.cuenta_descripcion,
            }
            for row in rows
        ],
        "total_debe": _decimal_to_float(total_debe),
        "total_haber": _decimal_to_float(total_haber),
    }


@erp_presupuesto_router.get("/panel/movimientos/export")
def export_erp_presupuesto_panel_movimientos(
    periodo: str = Query(..., description="Periodo en formato YYYY-MM"),
    proyecto_id: int = Query(..., description="Proyecto a consultar"),
    concepto: Literal["egreso", "ingreso", "resultado"] = Query(
        ..., description="Concepto real a descomponer"
    ),
    rubro_id: int | None = Query(None, description="Filtro opcional por rubro"),
    erp_cuenta_id: int | None = Query(None, description="Filtro opcional por cuenta"),
    session: Session = Depends(get_session),
):
    """Exporta a Excel los movimientos reales del panel."""
    try:
        anio, mes = _parse_month_key(periodo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    rows = session.exec(
        _build_panel_movimientos_stmt(
            anio,
            mes,
            proyecto_id,
            concepto,
            rubro_id,
            erp_cuenta_id,
        )
    ).all()
    total_debe = sum((row.debe or Decimal("0")) for row in rows)
    total_haber = sum((row.haber or Decimal("0")) for row in rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    headers = ["Fecha", "Asiento", "Cuenta", "Descripcion", "Debe", "Haber", "Neto"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="E5E7EB")
    header_font = Font(bold=True, color="111827")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        debe = row.debe or Decimal("0")
        haber = row.haber or Decimal("0")
        ws.append(
            [
                row.fecha,
                _format_movimiento_asiento(row),
                _format_movimiento_cuenta(row),
                row.descripcion or "",
                float(debe),
                float(haber),
                float(debe + haber),
            ]
        )

    ws.append(["", "", "", "Total", float(total_debe), float(total_haber), float(total_debe + total_haber)])
    total_row = ws.max_row
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F3F4F6")

    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[4].number_format = '$ #,##0.00'
        row[5].number_format = '$ #,##0.00'
        row[6].number_format = '$ #,##0.00'

    widths = {
        "A": 12,
        "B": 12,
        "C": 42,
        "D": 70,
        "E": 16,
        "F": 16,
        "G": 16,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"movimientos-reales-{concepto}-{periodo}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
        },
    )
